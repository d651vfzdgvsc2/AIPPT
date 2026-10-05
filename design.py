"""演示设计工作流核心：参考图提色 → 结构化设计规范 → 直接产出 .pptx。

设计原则（延续你另两个项目的气质）：
- 精确的（主色提取、版式渲染）交给算法；
- 理解的（风格语义、内容大纲）交给大模型；
- 规范可编辑、可复用（Skill 化）。
"""
from __future__ import annotations

import base64
import io
import json
import os
import re

import numpy as np
import requests

try:
    import cv2
except Exception:  # noqa: BLE001
    cv2 = None

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN


# ============================ 颜色 / 图像 ============================
def b64_to_bytes(image_base64: str) -> bytes:
    s = (image_base64 or "").strip()
    if s.startswith("data:") and "," in s:
        s = s.split(",", 1)[1]
    return base64.b64decode(s + "=" * (-len(s) % 4))


def _hex(rgb) -> str:
    return "#{:02X}{:02X}{:02X}".format(int(rgb[0]), int(rgb[1]), int(rgb[2]))


def _lum_hex(h: str) -> float:
    h = h.lstrip("#")
    r, g, b = (int(h[i:i + 2], 16) / 255 for i in (0, 2, 4))
    return 0.2126 * r + 0.7152 * g + 0.0722 * b


def _color_name(rgb) -> str:
    r, g, b = rgb
    mx, mn = max(rgb), min(rgb)
    if mx - mn < 28:
        if mx > 220:
            return "近白"
        if mx < 60:
            return "深灰/黑"
        return "中性灰"
    if r >= g and r >= b:
        return "暖红/橙" if g > b else "红/粉"
    if g >= r and g >= b:
        return "绿色系"
    return "蓝色系"


def extract_palette(image_base64: str, k: int = 5) -> list[dict]:
    """k-means 从参考图提取主色（确定性，不依赖模型）。"""
    if cv2 is None or not image_base64:
        return []
    try:
        raw = np.frombuffer(b64_to_bytes(image_base64), np.uint8)
        img = cv2.imdecode(raw, cv2.IMREAD_COLOR)
        if img is None:
            return []
        small = cv2.resize(img, (160, 160), interpolation=cv2.INTER_AREA)
        data = small.reshape(-1, 3).astype(np.float32)
        k = int(min(k, len(np.unique(data, axis=0))))
        if k < 2:
            return []
        criteria = (cv2.TERM_CRITERIA_EPS + cv2.TERM_CRITERIA_MAX_ITER, 20, 1.0)
        _, labels, centers = cv2.kmeans(data, k, None, criteria, 3, cv2.KMEANS_PP_CENTERS)
        counts = np.bincount(labels.flatten(), minlength=k)
        order = np.argsort(-counts)
        out = []
        for i in order:
            b, g, r = centers[i]
            rgb = (int(r), int(g), int(b))
            out.append({"hex": _hex(rgb), "name": _color_name(rgb),
                        "share": round(float(counts[i]) / counts.sum(), 2)})
        return out
    except Exception:  # noqa: BLE001
        return []


# ============================ 大模型调用 ============================
def _data_url(image_base64: str) -> str:
    s = (image_base64 or "").strip()
    return s if s.startswith("data:") else "data:image/png;base64," + s


def _chat(messages, api_url, api_key, model, max_tokens=1200, temperature=0.4) -> str:
    resp = requests.post(
        api_url,
        json={"model": model, "messages": messages,
              "max_tokens": max_tokens, "temperature": temperature},
        headers={"Authorization": f"Bearer {api_key}", "Content-Type": "application/json"},
        timeout=180,
    )
    resp.raise_for_status()
    data = resp.json()
    msg = data["choices"][0]["message"]
    return (msg.get("content") or msg.get("reasoning_content") or "").strip()


def _extract_json(text: str) -> dict:
    if not text:
        return {}
    m = re.search(r"\{[\s\S]*\}", text)
    if not m:
        return {}
    try:
        return json.loads(m.group(0))
    except ValueError:
        return {}


STYLE_SYSTEM = (
    "你是资深演示（PPT）设计师。根据用户给出的参考图与主色，输出**严格的 JSON**（不要任何多余文字）："
    '{"theme":"主题与气质(如：沉稳商务/科技极简)","layout":"版式特征(如：左文右图、大标题、三栏卡片)",'
    '"typography":"字体与层级(如：无衬线、标题特粗、正文常规)","mood_keywords":["关键词",...],'
    '"accent_usage":"强调色用法"}'
)

DECK_SYSTEM = (
    "你是资深演示设计师。根据用户需求与既定设计规范，输出**严格 JSON** 的演示大纲："
    '{"title":"主标题","subtitle":"副标题","slides":[{"title":"页标题","bullets":["要点1","要点2"]}]}。'
    "要求 6~10 页，每页 3~5 个要点，每个要点不超过 24 字，全部简体中文，不要输出任何解释文字。"
)


def analyze_style(image_base64, palette, api_url, api_key, model) -> dict:
    """用视觉模型把参考图理解成结构化设计规范（可编辑）。"""
    pal_txt = "、".join(f"{p['hex']}({p['name']})" for p in palette) or "无"
    content = [{"type": "text", "text": f"参考图主色：{pal_txt}。请分析这张演示参考图的设计风格。"}]
    if image_base64:
        content.append({"type": "image_url", "image_url": {"url": _data_url(image_base64)}})
    messages = [{"role": "system", "content": STYLE_SYSTEM},
                {"role": "user", "content": content}]
    try:
        spec = _extract_json(_chat(messages, api_url, api_key, model, max_tokens=800))
    except Exception:  # noqa: BLE001
        spec = {}
    spec.setdefault("theme", "")
    spec.setdefault("layout", "")
    spec.setdefault("typography", "")
    spec.setdefault("mood_keywords", [])
    spec.setdefault("accent_usage", "")
    return spec


def plan_deck(user_input, spec, api_url, api_key, model) -> dict:
    """根据需求 + 设计规范，生成演示大纲 JSON。"""
    spec_txt = json.dumps(spec, ensure_ascii=False)
    user = f"设计规范：{spec_txt}\n\n用户需求：{user_input}\n\n请据此输出演示大纲 JSON。"
    try:
        deck = _extract_json(_chat(
            [{"role": "system", "content": DECK_SYSTEM}, {"role": "user", "content": user}],
            api_url, api_key, model, max_tokens=2500))
    except Exception:  # noqa: BLE001
        deck = {}
    if not deck.get("slides"):
        deck = {"title": (user_input or "演示文稿")[:40], "subtitle": "",
                "slides": [{"title": "内容", "bullets": ["（大纲生成失败，请重试）"]}]}
    return deck


# ============================ 产出 .pptx ============================
def _rgb(h: str) -> RGBColor:
    h = (h or "#000000").lstrip("#")
    if len(h) != 6:
        h = "1F4468"
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def _role_colors(palette: list[dict]):
    """从主色里挑：背景(最亮) / 主色(最暗) / 强调色。"""
    hexes = [p["hex"] for p in (palette or []) if p.get("hex")] or ["#1F4468", "#F2F2F2"]
    ordered = sorted(hexes, key=_lum_hex)
    primary = ordered[0]
    bg = ordered[-1]
    if _lum_hex(bg) < 0.78:
        bg = "#FFFFFF"
    accent = next((h for h in ordered[1:] if abs(_lum_hex(h) - 0.45) < 0.3 and h != primary), primary)
    return bg, primary, accent


def build_pptx(spec: dict, deck: dict) -> bytes:
    """把设计规范 + 大纲渲染成 .pptx（16:9）。"""
    palette = spec.get("palette") or []
    bg, primary, accent = _role_colors(palette)
    text_c = "#1A1A1A"
    sub_c = "#5A5A5A"

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    def new_slide(bg_hex: str):
        s = prs.slides.add_slide(blank)
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = _rgb(bg_hex)
        return s

    def rect(slide, x, y, w, h, color):
        from pptx.enum.shapes import MSO_SHAPE
        shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
        shp.fill.solid()
        shp.fill.fore_color.rgb = _rgb(color)
        shp.line.fill.background()
        shp.shadow.inherit = False
        return shp

    def textbox(slide, x, y, w, h, runs):
        """runs: list of (text, size, bold, color, align)"""
        tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = True
        for i, (t, size, bold, color, align) in enumerate(runs):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = align
            r = p.add_run()
            r.text = t
            r.font.size = Pt(size)
            r.font.bold = bold
            r.font.color.rgb = _rgb(color)
            r.font.name = "Microsoft YaHei"
        return tb

    # 封面
    s = new_slide(bg)
    rect(s, 0.9, 2.45, 1.4, 0.12, accent)
    textbox(s, 0.9, 2.7, 11.6, 2.2, [
        (deck.get("title") or "演示文稿", 44, True, primary, PP_ALIGN.LEFT),
        (deck.get("subtitle") or "", 20, False, sub_c, PP_ALIGN.LEFT),
    ])
    textbox(s, 0.9, 6.6, 11.6, 0.6, [("由 DeckCraft · 参考图驱动的演示设计工作流生成", 12, False, sub_c, PP_ALIGN.LEFT)])

    # 内容页
    for slide_data in deck.get("slides", []):
        s = new_slide(bg)
        rect(s, 0.9, 1.15, 0.12, 0.75, accent)          # 左侧强调条
        textbox(s, 1.25, 1.05, 11.2, 1.0, [
            (slide_data.get("title") or "标题", 30, True, primary, PP_ALIGN.LEFT)])
        bullets = slide_data.get("bullets") or []
        if bullets:
            tb = s.shapes.add_textbox(Inches(1.25), Inches(2.35), Inches(11.0), Inches(4.5))
            tf = tb.text_frame
            tf.word_wrap = True
            for i, b in enumerate(bullets):
                p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                p.space_after = Pt(10)
                r = p.add_run()
                r.text = "•  " + str(b)
                r.font.size = Pt(19)
                r.font.color.rgb = _rgb(text_c)
                r.font.name = "Microsoft YaHei"

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()
